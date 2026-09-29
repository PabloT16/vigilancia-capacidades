"""Lectura de los ficheros de capacidad (CSV o XLSX) y aplanado de cabeceras multinivel.

REE publica con 3-4 filas de cabecera agrupadas (celdas combinadas); las distribuidoras
usan el formato normalizado de la CNMC con una sola fila. Aquí se convierten ambos casos
en una tabla con nombres de columna del tipo "GRUPO | Subgrupo | Columna".
"""
from __future__ import annotations

import csv
import io
import re
import unicodedata

import pandas as pd


def quitar_acentos(texto: str) -> str:
    s = unicodedata.normalize("NFKD", str(texto or ""))
    return "".join(c for c in s if not unicodedata.combining(c))


def norm(texto) -> str:
    """Texto normalizado para comparar: sin acentos, minúsculas y espacios simples."""
    return re.sub(r"\s+", " ", quitar_acentos(texto)).strip().lower()


def limpiar(texto) -> str:
    return re.sub(r"\s+", " ", str(texto or "")).strip()


def patron(expr: str) -> re.Pattern:
    """Compila una expresión del config.yaml ignorando mayúsculas y acentos."""
    return re.compile(quitar_acentos(expr), re.IGNORECASE)


# ---------------------------------------------------------------- números
_NUMERICO = re.compile(r"-?[\d.,]+\*?")
_MILES = re.compile(r"-?\d{1,3}(?:\.\d{3})+")


def detectar_decimal(celdas) -> str:
    """',' si el fichero usa coma decimal, '.' si usa punto decimal y 'miles' si el
    punto solo aparece como separador de miles (caso REE: 2.406 = 2406 MW)."""
    punto_decimal = False
    for v in celdas:
        s = str(v).strip().replace(" ", "")
        if not _NUMERICO.fullmatch(s):
            continue
        if "," in s:
            return ","
        if re.search(r"\.\d{1,2}\*?$|\.\d{4,}\*?$", s):
            punto_decimal = True
    return "." if punto_decimal else "miles"


def a_numero(valor, modo: str = "miles"):
    if valor is None:
        return None
    if isinstance(valor, (int, float)):
        return None if valor != valor else float(valor)
    s = str(valor).strip().replace("\xa0", "").replace(" ", "").rstrip("*")
    if s == "" or s.upper() in {"N/A", "NA", "-", "—", "ND", "NAN"}:
        return None
    if modo == ",":
        s = s.replace(".", "").replace(",", ".")
    elif modo == "miles":
        if "," in s:
            s = s.replace(".", "").replace(",", ".")
        elif _MILES.fullmatch(s):
            s = s.replace(".", "")
    try:
        return float(s)
    except ValueError:
        return None


# ---------------------------------------------------------------- lectura
def leer_filas(contenido: bytes, nombre: str = "") -> list[list[str]]:
    if nombre.lower().endswith((".xlsx", ".xlsm")) or contenido[:2] == b"PK":
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        return [["" if v is None else str(v) for v in fila] for fila in ws.iter_rows(values_only=True)]
    texto = None
    for codificacion in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            texto = contenido.decode(codificacion)
            break
        except UnicodeDecodeError:
            continue
    primera = texto.split("\n", 1)[0]
    separador = max([";", ",", "\t"], key=primera.count)
    return [fila for fila in csv.reader(io.StringIO(texto), delimiter=separador)]


def _parece_numero(celda: str) -> bool:
    return bool(_NUMERICO.fullmatch(celda.strip().replace(" ", "")))


def construir_tabla(filas: list[list[str]]) -> pd.DataFrame:
    """Localiza el bloque de cabecera, lo aplana y devuelve un DataFrame de textos."""
    inicio = next(i for i, f in enumerate(filas) if sum(1 for c in f if str(c).strip()) >= 3)
    fin = inicio + 1
    # Las filas de cabecera adicionales (REE) tienen la primera celda vacía y ningún número.
    while fin < len(filas) and fin - inicio < 6:
        fila = filas[fin]
        no_vacias = [str(c) for c in fila if str(c).strip()]
        if not fila or str(fila[0]).strip() or not no_vacias or any(_parece_numero(c) for c in no_vacias):
            break
        fin += 1
    bloque, datos = filas[inicio:fin], filas[fin:]
    ncol = max(len(f) for f in filas[inicio:])
    niveles = [[limpiar(c) for c in f] + [""] * (ncol - len(f)) for f in bloque]

    # Relleno hacia la derecha de las celdas combinadas; un nivel se reinicia cuando
    # empieza un grupo nuevo en cualquier nivel superior.
    rellenos = []
    for k, fila in enumerate(niveles):
        salida, ultimo = [], ""
        for c in range(ncol):
            if any(niveles[j][c] for j in range(k)):
                ultimo = ""
            if fila[c]:
                ultimo = fila[c]
            salida.append(ultimo)
        rellenos.append(salida)

    nombres, vistos = [], {}
    for c in range(ncol):
        partes = []
        for k in range(len(rellenos)):
            parte = rellenos[k][c]
            if parte and (not partes or partes[-1] != parte):
                partes.append(parte)
        nombre = " | ".join(partes) or f"columna {c + 1}"
        vistos[nombre] = vistos.get(nombre, 0) + 1
        if vistos[nombre] > 1:
            nombre = f"{nombre} #{vistos[nombre]}"
        nombres.append(nombre)

    registros = [
        [str(f[c]).strip() if c < len(f) else "" for c in range(ncol)]
        for f in datos
        if any(str(x).strip() for x in f)
    ]
    return pd.DataFrame(registros, columns=nombres, dtype=str)


def leer_tabla(contenido: bytes, nombre: str = "") -> pd.DataFrame:
    return construir_tabla(leer_filas(contenido, nombre))


# ---------------------------------------------------------------- selección de columnas
def buscar_columnas(df: pd.DataFrame, expr: str, excluir: str | None = None) -> list[str]:
    rx, ex = patron(expr), (patron(excluir) if excluir else None)
    return [c for c in df.columns if rx.search(norm(c)) and not (ex and ex.search(norm(c)))]


def primera_columna(df: pd.DataFrame, expr: str) -> str | None:
    columnas = buscar_columnas(df, expr)
    return columnas[0] if columnas else None


def seleccionar_metricas(df: pd.DataFrame, especificaciones: list[dict]) -> list[tuple[str, str, str]]:
    """Devuelve [(columna, etiqueta, categoría)] de las columnas numéricas a vigilar.

    - {grupo: regex}: todas las columnas cuyo grupo superior encaja; la etiqueta es el resto.
    - {columna: regex, excluir: regex}: columnas cuyo nombre completo encaja.
    """
    resultado, usadas = [], set()
    for spec in especificaciones:
        categoria = spec.get("categoria", "general")
        if "grupo" in spec:
            rx = patron(spec["grupo"])
            candidatas = [c for c in df.columns if rx.search(norm(c.split(" | ")[0]))]
        else:
            candidatas = buscar_columnas(df, spec["columna"], spec.get("excluir"))
        for col in candidatas:
            if col in usadas:
                continue
            valores = df[col].map(lambda v: a_numero(v, "miles"))
            if valores.notna().sum() == 0:
                continue  # columnas de relleno sin datos
            partes = col.split(" | ")
            etiqueta = " | ".join(partes[1:]) if "grupo" in spec and len(partes) > 1 else col
            resultado.append((col, re.sub(r"\s*\[MW\]", "", etiqueta).strip(), categoria))
            usadas.add(col)
    return resultado
