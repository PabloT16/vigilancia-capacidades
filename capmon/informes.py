"""Informe Excel de cada publicación y datos/HTML del panel publicado en GitHub Pages."""
from __future__ import annotations

import datetime as dt
import json
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .comparar import CATEGORIAS, Resultado

MADRID = ZoneInfo("Europe/Madrid")
FUENTE = "Arial"
HOJAS = {
    "afloramiento": "Afloramientos",
    "incremento": "Incrementos",
    "autoconsumo_afloramiento": "Autoconsumo aflora",
    "autoconsumo_incremento": "Autoconsumo sube",
    "reserva": "Reservas nuevas",
    "nudo_nuevo": "Nudos nuevos",
}


def ahora() -> dt.datetime:
    return dt.datetime.now(MADRID).replace(microsecond=0)


def fmt_fecha(d) -> str:
    if isinstance(d, str):
        d = dt.date.fromisoformat(d[:10])
    return d.strftime("%d/%m/%Y") if d else ""


# ---------------------------------------------------------------- Excel
def _cabecera(ws, titulos):
    relleno = PatternFill("solid", fgColor="1D2B36")
    for i, t in enumerate(titulos, 1):
        c = ws.cell(1, i, t)
        c.font = Font(name=FUENTE, bold=True, color="FFFFFF")
        c.fill = relleno
        c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"


def escribir_excel(ruta: Path, fuente: dict, actual, anterior, resultado: Resultado, umbral: float) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"
    filas = [
        ("Fuente", fuente["nombre"]),
        ("Publicación analizada", f"{fmt_fecha(actual.fecha)}  ({actual.url})"),
        ("Comparada con", f"{fmt_fecha(anterior.fecha)}  ({anterior.url})" if anterior else "Sin publicación anterior"),
        ("Nudos en la publicación", resultado.nudos_actual),
        ("Umbral de incremento", f"{umbral:g} MW"),
        ("Informe generado", ahora().strftime("%d/%m/%Y %H:%M")),
        ("", ""),
    ]
    cuenta = resultado.contadores()
    filas += [(CATEGORIAS[c], cuenta[c]) for c in CATEGORIAS if c in cuenta]
    filas += [("Aviso", a) for a in resultado.avisos]
    for fila in filas:
        ws.append(fila)
    for celda in ws["A"]:
        celda.font = Font(name=FUENTE, bold=True)
    for celda in ws["B"]:
        celda.font = Font(name=FUENTE)
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 110

    for categoria, titulo in CATEGORIAS.items():
        hallazgos = [h for h in resultado.hallazgos if h.categoria == categoria]
        if not hallazgos:
            continue
        hoja = wb.create_sheet(HOJAS[categoria])
        claves_ctx = []
        for h in hallazgos:
            claves_ctx += [k for k in h.contexto if k not in claves_ctx]
        if categoria == "reserva":
            titulos = ["Nudo", "Columna", "Valor actual y anterior"] + claves_ctx
        else:
            titulos = ["Nudo", "Capacidad", "Antes (MW)", "Ahora (MW)", "Variación (MW)"] + claves_ctx
        _cabecera(hoja, titulos)
        for h in hallazgos:
            base = [h.nudo, h.metrica, h.detalle] if categoria == "reserva" else [h.nudo, h.metrica, h.antes, h.ahora, h.delta]
            hoja.append(base + [h.contexto.get(k, "") for k in claves_ctx])
        for fila in hoja.iter_rows(min_row=2):
            for celda in fila:
                celda.font = Font(name=FUENTE)
                if isinstance(celda.value, float):
                    celda.number_format = "#,##0.00"
        anchos = [38, 60, 14, 14, 14] if categoria != "reserva" else [38, 42, 70]
        for i in range(1, len(titulos) + 1):
            hoja.column_dimensions[get_column_letter(i)].width = anchos[i - 1] if i <= len(anchos) else 22
        hoja.auto_filter.ref = hoja.dimensions
    ruta.parent.mkdir(parents=True, exist_ok=True)
    wb.save(ruta)
    return ruta


# ---------------------------------------------------------------- Panel
def _cargar(ruta: Path) -> dict:
    if ruta.exists():
        return json.loads(ruta.read_text(encoding="utf-8"))
    return {"fuentes": {}, "hallazgos": []}


def _base_fuente(fuente: dict) -> dict:
    return {k: fuente.get(k, "") for k in ("id", "nombre", "gestor", "tipo", "pagina")}


def registrar_publicacion(raiz: Path, fuente: dict, actual, anterior, resultado: Resultado, informe: Path,
                          sello: str) -> dict:
    ruta = raiz / "docs" / "data" / "panel.json"
    panel = _cargar(ruta)
    entrada = panel["fuentes"].get(fuente["id"], _base_fuente(fuente))
    entrada.update(_base_fuente(fuente))
    registro = {
        "fecha": actual.fecha.isoformat(),
        "url": actual.url,
        "sello": sello,
        "anterior": anterior.fecha.isoformat() if anterior else "",
        "detectado": ahora().isoformat(),
        "informe": informe.relative_to(raiz / "docs").as_posix(),
        "contadores": resultado.contadores(),
        "avisos": resultado.avisos,
    }
    entrada["ultima"] = registro
    entrada["error"] = ""
    entrada["historial"] = ([registro] + [h for h in entrada.get("historial", []) if h["url"] != actual.url])[:24]
    panel["fuentes"][fuente["id"]] = entrada
    panel["hallazgos"] = [h for h in panel["hallazgos"] if h["fuente_id"] != fuente["id"]] + [
        dict(h.dict(), fuente_id=fuente["id"], fuente=fuente["nombre"]) for h in resultado.hallazgos
    ]
    _guardar(ruta, panel)
    return registro


def registrar_error(raiz: Path, fuente: dict, mensaje: str) -> bool:
    """Guarda el error. Devuelve True si es un error nuevo (para no repetir el aviso en cada pasada)."""
    ruta = raiz / "docs" / "data" / "panel.json"
    panel = _cargar(ruta)
    entrada = panel["fuentes"].get(fuente["id"], _base_fuente(fuente))
    nuevo = entrada.get("error") != mensaje
    entrada["error"] = mensaje
    entrada["error_desde"] = entrada.get("error_desde") if not nuevo else ahora().isoformat()
    panel["fuentes"][fuente["id"]] = entrada
    _guardar(ruta, panel)
    return nuevo


def limpiar_error(raiz: Path, fuente: dict) -> None:
    ruta = raiz / "docs" / "data" / "panel.json"
    panel = _cargar(ruta)
    if fuente["id"] in panel["fuentes"] and panel["fuentes"][fuente["id"]].get("error"):
        panel["fuentes"][fuente["id"]]["error"] = ""
        _guardar(ruta, panel)


def _guardar(ruta: Path, panel: dict) -> None:
    ruta.parent.mkdir(parents=True, exist_ok=True)
    ruta.write_text(json.dumps(panel, ensure_ascii=False, indent=1), encoding="utf-8")


def generar_html(raiz: Path, config: dict) -> Path:
    panel = _cargar(raiz / "docs" / "data" / "panel.json")
    orden = [f["id"] for f in config["fuentes"]]
    datos = {
        "actualizado": ahora().isoformat(),
        "umbral": config.get("umbral_incremento_mw", 10),
        "fuentes": [panel["fuentes"].get(i, {"id": i, "nombre": next(f["nombre"] for f in config["fuentes"] if f["id"] == i)})
                    for i in orden],
        "hallazgos": panel["hallazgos"],
        "categorias": CATEGORIAS,
    }
    plantilla = (Path(__file__).parent / "panel.html").read_text(encoding="utf-8")
    json_seguro = json.dumps(datos, ensure_ascii=False).replace("</", "<\\/")
    salida = raiz / "docs" / "index.html"
    salida.write_text(plantilla.replace("__DATOS__", json_seguro), encoding="utf-8")
    return salida
