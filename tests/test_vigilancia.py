import datetime as dt
import json
from pathlib import Path

import pytest
import yaml
from openpyxl import load_workbook

from capmon import __main__ as app
from capmon import fuentes as F
from capmon.avisos import componer
from capmon.comparar import comparar
from capmon.informes import escribir_excel
from capmon.tabla import a_numero, detectar_decimal, leer_tabla, seleccionar_metricas

RAIZ = Path(__file__).resolve().parents[1]
FIX = RAIZ / "tests" / "fixtures"
CONFIG = yaml.safe_load((RAIZ / "config.yaml").read_text(encoding="utf-8"))
FUENTE = {f["id"]: f for f in CONFIG["fuentes"]}


def tabla(nombre, bom=False):
    datos = (FIX / nombre).read_bytes()
    return leer_tabla(("\ufeff".encode() + datos) if bom else datos, nombre)


def columna(df, fuente_id, etiqueta):
    for col, et, _ in seleccionar_metricas(df, FUENTE[fuente_id]["metricas"]):
        if et == etiqueta:
            return col
    raise KeyError(etiqueta)


def fila(df, clave):
    return df.index[df.iloc[:, 0].str.strip() == clave][0]


# ---------------------------------------------------------------- números
@pytest.mark.parametrize("texto,modo,esperado", [
    ("2.406", "miles", 2406), ("1.209", "miles", 1209), ("428", "miles", 428), ("27,43", ",", 27.43),
    ("2 *", "miles", 2), ("N/A", "miles", None), ("", ",", None), ("0,05", ",", 0.05), ("12.5", ".", 12.5)])
def test_a_numero(texto, modo, esperado):
    assert a_numero(texto, modo) == esperado


def test_detectar_decimal():
    assert detectar_decimal(["2.406", "428", "0"]) == "miles"
    assert detectar_decimal(["27,43", "0"]) == ","
    assert detectar_decimal(["12.5", "3"]) == "."


# ---------------------------------------------------------------- REE generación
def test_ree_generacion_columnas_reales():
    df = tabla("ree_generacion.csv", bom=True)
    assert df.columns[0] == "Nombre y tensión del nudo"
    etiquetas = [(et, cat) for _, et, cat in seleccionar_metricas(df, FUENTE["ree_generacion"]["metricas"])]
    assert ("Generación | MPE RdT", "general") in etiquetas
    assert ("Almacenamiento | MPE RdD", "general") in etiquetas
    assert sum(1 for _, c in etiquetas if c == "autoconsumo") == 6
    assert len(etiquetas) == 14


def test_ree_generacion_deteccion():
    act = tabla("ree_generacion.csv")
    ant = act.copy()
    ant.loc[fila(ant, "ALCOCERO DE MOLA 220"), columna(ant, "ree_generacion", "Generación | MPE RdT")] = "0"
    ant.loc[fila(ant, "ALCOCERO DE MOLA 220"),
            columna(ant, "ree_generacion", "Generación | MGES RdTo MPE con compensador síncrono RdT")] = "130"
    ant.loc[fila(ant, "ABADES 400"), columna(ant, "ree_generacion", "Capacidad de acceso liberada para MPE")] = "0"
    ant.loc[fila(ant, "ABADES 400"), columna(ant, "ree_generacion", "Capacidad de acceso liberada para MGES")] = "40"
    ant.loc[fila(ant, "GAROÑA 220"), "Motivo de reserva de capacidad de acceso en el nudo"] = ""

    res = comparar(ant, act, FUENTE["ree_generacion"])
    por = {(h.categoria, h.nudo, h.metrica): h for h in res.hallazgos}
    assert ("afloramiento", "ALCOCERO DE MOLA 220", "Generación | MPE RdT") in por
    inc = por[("incremento", "ALCOCERO DE MOLA 220", "Generación | MGES RdTo MPE con compensador síncrono RdT")]
    assert (inc.antes, inc.ahora, inc.delta) == (130, 145, 15)
    assert ("autoconsumo_afloramiento", "ABADES 400", "Capacidad de acceso liberada para MPE") in por
    assert por[("autoconsumo_incremento", "ABADES 400", "Capacidad de acceso liberada para MGES")].delta == 15
    reserva = por[("reserva", "GAROÑA 220", "Motivo de reserva de capacidad de acceso en el nudo")]
    assert reserva.detalle.startswith("NTJ")
    assert reserva.contexto["Comunidad Autónoma"] == "Castilla y León"
    assert res.contadores()["afloramiento"] == 1


def test_umbral_no_salta_por_debajo():
    act = tabla("ree_generacion.csv")
    ant = act.copy()
    ant.loc[fila(ant, "ALCOCERO DE MOLA 220"),
            columna(ant, "ree_generacion", "Generación | MGES RdTo MPE con compensador síncrono RdT")] = "136"
    assert not [h for h in comparar(ant, act, FUENTE["ree_generacion"]).hallazgos if h.categoria == "incremento"]


# ---------------------------------------------------------------- REE demanda
def test_ree_demanda_deteccion():
    act = tabla("ree_demanda.csv")
    assert [et for _, et, _ in seleccionar_metricas(act, FUENTE["ree_demanda"]["metricas"])] == [
        "DEMANDA | CEP | CH", "DEMANDA | CEP | SH", "DEMANDA | NO CEP", "ALMACENAMIENTO | CEP (MPE)", "ALMACENAMIENTO | NO CEP (MGES)"]
    ant = act.copy()
    ant.loc[fila(ant, "ABADIANO 220"), columna(ant, "ree_demanda", "ALMACENAMIENTO | CEP (MPE)")] = ""  # vacío = 0
    ant.loc[fila(ant, "ALGECIRAS 220"), "Nudo de concurso"] = ""
    res = comparar(ant, act, FUENTE["ree_demanda"])
    categorias = {(h.categoria, h.nudo) for h in res.hallazgos}
    assert ("afloramiento", "ABADIANO 220") in categorias
    assert ("reserva", "ALGECIRAS 220") in categorias


# ---------------------------------------------------------------- i-DE (formato CNMC)
def test_ide_decimales_y_nudo_nuevo():
    act = tabla("ide_generacion.csv")
    col = columna(act, "ide_generacion", "Capacidad disponible (MW)")
    ant = act.copy()
    clave = "Identificador del Punto de Conexión"
    ant.loc[ant[clave] == "0081004267", col] = "150,5"   # BENIDORM +10,56
    ant.loc[ant[clave] == "0403065029", col] = "0"       # CALPE aflora
    ant.loc[ant[clave] == "0031001374", col] = "10"      # MIRANDA +6,36: no llega al umbral
    ant = ant[ant[clave] != "0025230000"]                # ALI no existía
    res = comparar(ant, act, FUENTE["ide_generacion"])
    por = {(h.categoria, h.nudo): h for h in res.hallazgos}
    assert por[("incremento", "BENIDORM 132.00")].delta == pytest.approx(10.56)
    assert por[("afloramiento", "CALPE T1")].ahora == pytest.approx(27.91)
    assert por[("nudo_nuevo", "ST STR ALI A30.000")].ahora == pytest.approx(27.43)
    assert ("incremento", "MIRANDA T5 30") not in por
    assert por[("incremento", "BENIDORM 132.00")].contexto["Nudo Afección RdT"] == "JIJONA 220"


# ---------------------------------------------------------------- descubrimiento de enlaces
HTML_REE = """<ul>
<li><a href="https://www.ree.es/sites/default/files/12_CLIENTES/Documentos/2026_09_01_GRT_generacion.pdf">pdf</a></li>
<li><a href="https://www.ree.es/sites/default/files/12_CLIENTES/Documentos/2026_09_01_GRT_generacion.csv">csv</a></li>
<li><a href="https://www.ree.es/sites/default/files/12_CLIENTES/Documentos/2026_09_01_GRT_generacion.xlsx">xlsx</a></li>
<li><a href="https://www.ree.es/sites/default/files/12_CLIENTES/Documentos/2026_09_01_Generacion_por_posicion.pdf">pos</a></li>
</ul><p>Información actualizada a 1 de septiembre de 2026:</p>"""

HTML_IDE = """<p>También puedes descargar:
<a href="https://www.i-de.es/documents/d/guest/2026_09_04_r1-001_generacion"><img src="https://www.i-de.es/documents/d/guest/pdf-svg"></a> /
<a href="https://www.i-de.es/documents/d/guest/2026_09_04_r1-001_generacion-1-"><img src="https://www.i-de.es/documents/d/guest/csv-svg"></a> /
<a href="https://www.i-de.es/documents/d/guest/2026_09_04_r1-001_generacion-2-"><img src="https://www.i-de.es/documents/d/guest/xlsx-svg"></a></p>
<p>Día y hora de publicación del PDF: 04/09/2026 a las 13:30H. </p>
<h3>2026</h3> 3 de agosto:
<a href="https://www.i-de.es/documents/d/guest/2026_08_03_r1-001_generacion-2-"><img src="https://www.i-de.es/documents/d/guest/pdf-svg"></a> /
<a href="https://www.i-de.es/documents/d/guest/2026_08_03_r1-001_generacion-1-"><img src="https://www.i-de.es/documents/d/guest/csv-svg"></a>"""

HTML_EDIS = """<a href="https://www.edistribucion.com/content/dam/edistribucion/conexion-a-la-red/descargables/nodos/generacion/202609/2026_09_03_R1299_generación.csv">Capacidad de generación en e-distribución septiembre 2026 csv</a>
<a href="https://www.edistribucion.com/content/dam/edistribucion/conexion-a-la-red/descargables/nodos/generacion/202609/2026_09_03_R1026_generaci%C3%B3n.csv">EASA</a>
<a href="/content/dam/edistribucion/conexion-a-la-red/descargables/nodos/generacion/202608/2026_08_03_R1299_generaci%C3%B3n.csv">agosto</a>"""


def test_descubrir_ree():
    enlaces = F.descubrir(HTML_REE, "https://www.ree.es/", FUENTE["ree_generacion"]["patron_fichero"])
    elegido = F.elegir(enlaces)
    assert elegido.fecha == dt.date(2026, 9, 1) and elegido.formato == "csv"
    assert F.sello(HTML_REE, FUENTE["ree_generacion"]["sello"]) == "1 de septiembre de 2026"


def test_descubrir_ide_por_icono():
    enlaces = F.descubrir(HTML_IDE, "https://www.i-de.es/", FUENTE["ide_generacion"]["patron_fichero"])
    actual = F.elegir(enlaces)
    assert actual.url.endswith("2026_09_04_r1-001_generacion-1-") and actual.formato == "csv"
    previo = F.elegir(enlaces, antes_de=actual.fecha)
    assert previo.fecha == dt.date(2026, 8, 3) and previo.url.endswith("-1-")
    assert F.sello(HTML_IDE, FUENTE["ide_generacion"]["sello"]) == "04/09/2026 a las 13:30H"


def test_descubrir_edistribucion_acentos_y_relativos():
    base = "https://www.edistribucion.com/es/red-electrica/nodos-capacidad-red/capacidad-generacion.html"
    enlaces = F.descubrir(HTML_EDIS, base, FUENTE["edistribucion_generacion"]["patron_fichero"])
    assert [e.fecha for e in enlaces] == [dt.date(2026, 9, 3), dt.date(2026, 8, 3)]   # EASA (R1026) excluida
    assert enlaces[1].url.startswith("https://www.edistribucion.com/content/")


# ---------------------------------------------------------------- informes y avisos
def test_excel_y_mensaje(tmp_path):
    act = tabla("ide_generacion.csv")
    ant = act.copy()
    ant.loc[ant["Identificador del Punto de Conexión"] == "0403065029", "Capacidad disponible (MW)"] = "0"
    res = comparar(ant, act, FUENTE["ide_generacion"])
    e_act = F.Enlace(dt.date(2026, 9, 4), "csv", "https://x/2026_09_04.csv")
    e_ant = F.Enlace(dt.date(2026, 8, 3), "csv", "https://x/2026_08_03.csv")
    ruta = escribir_excel(tmp_path / "i.xlsx", FUENTE["ide_generacion"], e_act, e_ant, res, 10)
    wb = load_workbook(ruta)
    assert wb.sheetnames == ["Resumen", "Afloramientos"]
    assert wb["Afloramientos"]["A2"].value == "CALPE T1" and wb["Afloramientos"]["D2"].value == pytest.approx(27.91)
    asunto, cuerpo = componer(FUENTE["ide_generacion"], e_act, e_ant, res, "https://panel/informe.xlsx")
    assert "04/09/2026" in asunto and "CALPE T1" in cuerpo and "0 → 27,91 MW" in cuerpo


# ---------------------------------------------------------------- flujo completo sin red
class Respuesta:
    def __init__(self, texto, url):
        self.text, self.url, self.status_code = texto, url, 200

    def raise_for_status(self):
        pass


def test_flujo_completo(tmp_path, monkeypatch):
    raiz = tmp_path / "repo"
    raiz.mkdir()
    cfg = dict(CONFIG, fuentes=[FUENTE["ide_generacion"]])
    (raiz / "config.yaml").write_text(yaml.safe_dump(cfg, allow_unicode=True), encoding="utf-8")
    septiembre = (FIX / "ide_generacion.csv").read_bytes()
    agosto = septiembre.replace(b";3198;20;27,91;", b";3198;20;0;")
    ficheros = {"2026_09_04_r1-001_generacion-1-": septiembre, "2026_08_03_r1-001_generacion-1-": agosto}
    pagina = {"html": HTML_IDE}

    monkeypatch.setattr(F, "pedir", lambda s, url, *a, **k: Respuesta(pagina["html"], url))
    monkeypatch.setattr(F, "firma", lambda s, url: "etag-" + url[-30:] + str(len(ficheros[url.rsplit("/", 1)[1]])))
    monkeypatch.setattr(F, "descargar", lambda s, url: ficheros[url.rsplit("/", 1)[1]])
    enviados = []
    monkeypatch.setattr("capmon.avisos.enviar", lambda asunto, cuerpo, adjunto=None: enviados.append(asunto) or [])

    # 1) arranque: compara septiembre con agosto sacado del histórico de la propia página
    assert app.vigilar(raiz) == 0
    estado = json.loads((raiz / "state" / "estado.json").read_text())
    assert estado["ide_generacion"]["fecha"] == "2026-09-04"
    panel = json.loads((raiz / "docs" / "data" / "panel.json").read_text())
    assert [h["nudo"] for h in panel["hallazgos"] if h["categoria"] == "afloramiento"] == ["CALPE T1"]
    assert len(enviados) == 1 and "arranque inicial" in enviados[0]

    # 2) nada cambia: no descarga ni avisa
    assert app.vigilar(raiz) == 0 and len(enviados) == 1

    # 3) llega octubre con 25 MW más en BENIDORM 132
    octubre = septiembre.replace(b";132;161,06;", b";132;186,06;")
    ficheros["2026_10_02_r1-001_generacion-1-"] = octubre
    pagina["html"] = HTML_IDE.replace("<p>También", '<a href="https://www.i-de.es/documents/d/guest/'
                                      '2026_10_02_r1-001_generacion-1-"><img src="csv-svg"></a><p>También')
    assert app.vigilar(raiz) == 0
    panel = json.loads((raiz / "docs" / "data" / "panel.json").read_text())
    incrementos = [h for h in panel["hallazgos"] if h["categoria"] == "incremento"]
    assert [(h["nudo"], h["delta"]) for h in incrementos] == [("BENIDORM 132.00", 25.0)]
    assert len(enviados) == 2 and "02/10/2026" in enviados[1]
    assert (raiz / "docs" / "informes" / "ide_generacion_2026-10-02.xlsx").exists()
    html = (raiz / "docs" / "index.html").read_text(encoding="utf-8")
    assert "BENIDORM 132.00" in html and "__DATOS__" not in html


def test_config_sin_claves_duplicadas():
    """Una clave repetida en YAML pisa a la anterior sin avisar (p. ej. 'nombre')."""
    class Estricto(yaml.SafeLoader):
        pass

    def mapa(loader, nodo, deep=False):
        claves = [loader.construct_object(k, deep=deep) for k, _ in nodo.value]
        assert len(claves) == len(set(claves)), f"clave duplicada en config.yaml: {claves}"
        return loader.construct_mapping(nodo, deep)

    Estricto.add_constructor(yaml.resolver.BaseResolver.DEFAULT_MAPPING_TAG, mapa)
    datos = yaml.load((RAIZ / "config.yaml").read_text(encoding="utf-8"), Loader=Estricto)
    assert [f["nombre"] for f in datos["fuentes"]] == [
        "REE generación", "REE demanda", "i-DE generación", "i-DE demanda", "e-distribución generación", "e-distribución demanda"]
